"""Report generation from frozen snapshots with SHA-256 checksums, and XLSX / PDF / CSV downloads."""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import uuid
from typing import Literal

from fastapi import APIRouter, Depends, Query
from fastapi.responses import Response
from pydantic import BaseModel
from sqlalchemy.orm import Session

from app.core.db import get_db
from app.core.errors import invalid, not_found
from app.models import ReportSnapshot
from app.modules.audit import chain
from app.modules.auth.deps import current_user
from app.modules.dashboard.router import _scope
from app.modules.policy.roles import REPORT_ROLES, CurrentUser, Role
from app.modules.reports.catalogue import CATALOGUE, period_bounds
from app.modules.sla.calendar import today_ist

router = APIRouter(prefix="/reports", tags=["reports"])


def checksum(report_code: str, params: dict, scope: dict, cut_off: dt.datetime, data: dict) -> str:
    cut = cut_off.astimezone(dt.UTC).isoformat()
    canonical = json.dumps(
        {"report": report_code, "params": params, "scope": scope, "cut_off": cut, "data": data},
        sort_keys=True,
        separators=(",", ":"),
        ensure_ascii=False,
        default=str,
    )
    return hashlib.sha256(canonical.encode()).hexdigest()


@router.get("/catalogue")
def catalogue(user: CurrentUser = Depends(current_user)) -> dict:
    return {
        "items": [
            {"code": s.code, "title": s.title, "description": s.description, "period": s.period}
            for s in CATALOGUE.values()
        ]
    }


class GenerateIn(BaseModel):
    report_code: str
    unit: str | None = None
    region: int | None = None
    period: Literal["month", "quarter"] | None = None
    anchor: dt.date | None = None


@router.post("", status_code=201)
def generate(body: GenerateIn, user: CurrentUser = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    user.require(REPORT_ROLES, "generate reports")
    spec = CATALOGUE.get(body.report_code)
    if spec is None:
        raise invalid("Unknown report.")
    if user.has(Role.AUDITOR) and not user.sees_all_units:
        raise invalid("Auditors generate organisation-wide reports only.")
    where, p, scope = _scope(db, user, body.unit, body.region)
    today = today_ist()
    start, end, label = period_bounds(body.period, body.anchor, today)
    opts = {"today": today, "start": start, "end": end}
    cut_off = dt.datetime.now(dt.UTC).replace(microsecond=0)
    data = spec.builder(db, where, p, opts)
    params = {"period": body.period, "anchor": body.anchor.isoformat() if body.anchor else None}
    if spec.period:
        data["period"] = {"from": start.isoformat(), "to": end.isoformat(), "label": label}
    data = json.loads(json.dumps(data, default=str))
    title = spec.title + (f" — {label}" if spec.period else "")
    digest = checksum(spec.code, params, scope, cut_off, data)
    snap = ReportSnapshot(
        id=uuid.uuid4(),
        report_code=spec.code,
        title=title,
        params=params,
        scope=scope,
        cut_off=cut_off,
        data=data,
        checksum=digest,
        created_by=user.id,
    )
    db.add(snap)
    chain.record(
        db,
        "report.generated",
        actor_id=user.id,
        actor_unit_id=user.unit_id,
        entity_type="report",
        entity_id=snap.id,
        payload={"report": spec.code, "checksum": digest, "scope": scope},
    )
    db.commit()
    return _snapshot_payload(snap, include_data=True)


def _snapshot_payload(s: ReportSnapshot, include_data: bool = False) -> dict:
    out = {
        "id": str(s.id),
        "report_code": s.report_code,
        "title": s.title,
        "params": s.params,
        "scope": s.scope,
        "cut_off": s.cut_off.isoformat(),
        "checksum": s.checksum,
        "created_at": s.created_at.isoformat(),
    }
    if include_data:
        out["data"] = s.data
    return out


@router.get("")
def list_snapshots(
    limit: int = Query(default=30, ge=1, le=100),
    user: CurrentUser = Depends(current_user),
    db: Session = Depends(get_db),
) -> dict:
    user.require(REPORT_ROLES, "view reports")
    q = db.query(ReportSnapshot)
    if not user.sees_all_units:
        q = q.filter(ReportSnapshot.created_by == user.id)
    snaps = q.order_by(ReportSnapshot.created_at.desc()).limit(limit).all()
    return {"items": [_snapshot_payload(s) for s in snaps]}


def _load(db: Session, user: CurrentUser, sid: uuid.UUID) -> ReportSnapshot:
    snap = db.get(ReportSnapshot, sid)
    if snap is None or (not user.sees_all_units and snap.created_by != user.id):
        raise not_found("Report")
    return snap


@router.get("/{sid}")
def get_snapshot(sid: uuid.UUID, user: CurrentUser = Depends(current_user), db: Session = Depends(get_db)) -> dict:
    user.require(REPORT_ROLES, "view reports")
    snap = _load(db, user, sid)
    recomputed = checksum(snap.report_code, snap.params, snap.scope, snap.cut_off, snap.data)
    return {**_snapshot_payload(snap, include_data=True), "checksum_verified": recomputed == snap.checksum}


@router.get("/{sid}/download")
def download(
    sid: uuid.UUID,
    fmt: Literal["xlsx", "csv", "pdf"] = "xlsx",
    user: CurrentUser = Depends(current_user),
    db: Session = Depends(get_db),
) -> Response:
    user.require(REPORT_ROLES, "download reports")
    snap = _load(db, user, sid)
    name = f"{snap.report_code.lower()}-{snap.cut_off.strftime('%Y%m%d-%H%M')}"
    if fmt == "csv":
        content, mime = _csv(snap), "text/csv; charset=utf-8"
    elif fmt == "pdf":
        content, mime = _pdf(snap), "application/pdf"
    else:
        content, mime = _xlsx(snap), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    chain.record(
        db,
        "report.downloaded",
        actor_id=user.id,
        actor_unit_id=user.unit_id,
        entity_type="report",
        entity_id=snap.id,
        payload={"format": fmt, "checksum": snap.checksum},
    )
    db.commit()
    return Response(
        content=content,
        media_type=mime,
        headers={"Content-Disposition": f'attachment; filename="{name}.{fmt}"', "X-Report-Checksum": snap.checksum},
    )


def _footer(s: ReportSnapshot) -> list[str]:
    return [
        f"Cut-off: {s.cut_off.isoformat()}",
        f"SHA-256: {s.checksum}",
        "Generated by Samanvay from a frozen snapshot; re-downloading gives identical figures.",
    ]


def _csv(s: ReportSnapshot) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([s.title])
    w.writerow(s.data["columns"])
    w.writerows(s.data["rows"])
    if s.data.get("totals"):
        w.writerow(s.data["totals"])
    w.writerow([])
    for line in [*s.data.get("notes", []), *_footer(s)]:
        w.writerow([line])
    return ("﻿" + buf.getvalue()).encode()


def _xlsx(s: ReportSnapshot) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = s.report_code[:31]
    ws.append([s.title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(s.data["columns"])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for row in s.data["rows"]:
        ws.append(row)
    if s.data.get("totals"):
        ws.append(s.data["totals"])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    ws.append([])
    for line in [*s.data.get("notes", []), *_footer(s)]:
        ws.append([line])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(
            48, max(10, *(len(str(c.value or "")) + 2 for c in col[1:]))
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(s: ReportSnapshot) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), title=s.title, author="Samanvay")
    styles = getSampleStyleSheet()
    rows = [s.data["columns"], *s.data["rows"]]
    if s.data.get("totals"):
        rows.append(s.data["totals"])
    table = Table([[("" if v is None else str(v)) for v in r] for r in rows], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#262B5E")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9CCE0")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F5FA")]),
    ]
    if s.data.get("totals"):
        style.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    story = [Paragraph(s.title, styles["Title"]), Spacer(1, 6), table, Spacer(1, 10)]
    for line in [*s.data.get("notes", []), *_footer(s)]:
        story.append(Paragraph(line, styles["Normal"]))
    doc.build(story)
    return out.getvalue()
