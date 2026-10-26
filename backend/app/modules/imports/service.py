"""Register Importer (blueprint SOL-01, W7): any unit's spreadsheet or a CPGRAMS export → canonical cases.

Column and status mappings are suggested from header names and values with string similarity (no AI needed);
rows are validated one by one inside savepoints so a bad row is reported and skipped, never half-written; re-running
the same file is idempotent on (channel, source_ref).
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import logging
import re
import uuid
from pathlib import Path
from typing import Any

from rapidfuzz import fuzz, process
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.models import ImportJob
from app.modules.case import fsm
from app.modules.case.service import CHANNELS, CaseInput, create_case
from app.modules.documents.storage import store
from app.modules.policy.roles import CurrentUser

log = logging.getLogger("samanvay.imports")

FIELDS: dict[str, list[str]] = {
    "received_on": [
        "date of receipt",
        "receipt date",
        "received on",
        "date received",
        "received",
        "diary date",
        "registration date",
        "date",
    ],
    "institution_name": [
        "institution",
        "institute",
        "college",
        "name of institution",
        "institution name",
        "name of the institute",
        "institute name",
    ],
    "permanent_id": ["permanent id", "aicte id", "pid", "aicte permanent id", "institute id"],
    "category": ["category", "nature of complaint", "complaint type", "grievance category", "type"],
    "subject": ["subject", "title", "summary"],
    "body": [
        "details",
        "description",
        "complaint",
        "grievance text",
        "complaint details",
        "brief",
        "grievance",
        "gist of complaint",
        "gist",
        "grievance description",
    ],
    "source_ref": [
        "registration number",
        "registration no",
        "reg no",
        "reference no",
        "ref no",
        "cpgrams no",
        "grievance id",
        "diary no",
        "reference",
    ],
    "unit_file_no": ["file no", "file number", "eoffice file", "e-office file no"],
    "status": ["status", "current status", "stage"],
    "complainant_name": ["complainant", "name of complainant", "complainant name", "petitioner", "name"],
    "complainant_phone": ["mobile", "phone", "contact no", "mobile no", "phone number"],
    "complainant_email": ["email", "e-mail", "email id"],
    "complainant_type": ["complainant type", "category of complainant"],
    "city": ["city", "district", "place", "location"],
    "state": ["state"],
    "pin": ["pin", "pincode", "pin code"],
    "disposed_on": ["date of disposal", "disposed on", "closure date", "disposal date"],
    "next_action": ["next action", "action to be taken"],
}
REQUIRED = ("received_on", "category", "body")

STATUS_HINTS = {
    "received": "RECEIVED",
    "new": "RECEIVED",
    "fresh": "RECEIVED",
    "pending": "REGISTERED",
    "registered": "REGISTERED",
    "open": "REGISTERED",
    "under process": "UNDER_EXAMINATION",
    "under examination": "UNDER_EXAMINATION",
    "in process": "UNDER_EXAMINATION",
    "processing": "UNDER_EXAMINATION",
    "under consideration": "UNDER_EXAMINATION",
    "reply awaited": "AWAITING_INFORMATION",
    "awaiting reply": "AWAITING_INFORMATION",
    "comments sought": "AWAITING_INFORMATION",
    "information sought": "AWAITING_INFORMATION",
    "action initiated": "ACTION_INITIATED",
    "show cause issued": "ACTION_INITIATED",
    "scn issued": "ACTION_INITIATED",
    "committee visit": "ACTION_INITIATED",
    "hearing": "HEARING",
    "hearing scheduled": "HEARING",
    "decision taken": "DECISION_RECORDED",
    "disposed": "DISPOSED",
    "closed": "DISPOSED",
    "resolved": "DISPOSED",
    "disposed of": "DISPOSED",
    "rejected": "REJECTED_AT_INTAKE",
    "not pertaining": "REJECTED_AT_INTAKE",
}
DATE_FORMATS = ["%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y-%m-%d", "%d-%b-%Y", "%d %b %Y", "%d-%m-%y", "%d/%m/%y"]
MAX_ROWS = 20000
MAX_ERRORS = 1000


def read_table(data: bytes, filename: str) -> tuple[list[str], list[list[Any]]]:
    ext = Path(filename).suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif ext == ".csv":
        textdata = data.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(textdata)))
    else:
        raise AppError(415, "File type not allowed", "Upload a .csv or .xlsx register.", "unsupported-type")
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise AppError(422, "Empty file", "The file has no rows.", "empty")
    header = [str(h or "").strip() for h in rows[0]]
    if len(rows) - 1 > MAX_ROWS:
        raise AppError(413, "Too many rows", f"Split files larger than {MAX_ROWS} rows.", "too-large")
    return header, rows[1:]


def _clean(h: str) -> str:
    return re.sub(r"[^a-z0-9 ]", " ", h.lower()).strip()


def suggest_columns(header: list[str]) -> dict[str, str | None]:
    """Assign each canonical field its best-matching header, best pairs first, each header used once."""
    choices = {_clean(h): h for h in header if h}
    scored = []
    for fld, synonyms in FIELDS.items():
        for clean, original in choices.items():
            best = max(0.5 * fuzz.token_set_ratio(syn, clean) + 0.5 * fuzz.ratio(syn, clean) for syn in synonyms)
            scored.append((best, fld, original))
    mapping: dict[str, str | None] = dict.fromkeys(FIELDS)
    used: set[str] = set()
    for score, fld, original in sorted(scored, key=lambda t: -t[0]):
        if score < 75 or mapping[fld] or original in used:
            continue
        mapping[fld] = original
        used.add(original)
    return mapping


def suggest_status(value: str) -> str | None:
    v = _clean(str(value))
    if not v:
        return None
    if v.upper().replace(" ", "_") in fsm.STATUSES:
        return v.upper().replace(" ", "_")
    hit = process.extractOne(v, list(STATUS_HINTS), scorer=fuzz.token_set_ratio)
    return STATUS_HINTS[hit[0]] if hit and hit[1] >= 80 else None


def suggest_category(db: Session, value: str) -> str | None:
    rows = db.execute(text("SELECT code, name FROM category WHERE is_active")).all()
    v = str(value).strip()
    for r in rows:
        if v.upper() == r.code:
            return r.code
    hit = process.extractOne(v.lower(), {r.code: r.name.lower() for r in rows}, scorer=fuzz.token_set_ratio)
    return hit[2] if hit and hit[1] >= 70 else None


def preview(db: Session, data: bytes, filename: str) -> dict:
    header, rows = read_table(data, filename)
    cols = suggest_columns(header)
    idx = {h: i for i, h in enumerate(header)}

    def distinct(fld: str) -> list[str]:
        col = cols.get(fld)
        if not col:
            return []
        vals = {str(r[idx[col]]).strip() for r in rows if idx[col] < len(r) and r[idx[col]] not in (None, "")}
        return sorted(vals)[:50]

    return {
        "header": header,
        "row_count": len(rows),
        "sample": [[("" if c is None else str(c)) for c in r] for r in rows[:10]],
        "column_map": cols,
        "status_map": {v: suggest_status(v) for v in distinct("status")},
        "category_map": {v: suggest_category(db, v) for v in distinct("category")},
        "required": list(REQUIRED),
        "fields": list(FIELDS),
        "statuses": fsm.STATUSES,
        "channels": CHANNELS,
    }


def _parse_date(value: Any, fmt: str | None) -> dt.date | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()
    for f in ([fmt] if fmt else []) + DATE_FORMATS:
        try:
            return dt.datetime.strptime(s, f).date()
        except ValueError:
            continue
    return None


def process_import(db: Session, job: ImportJob, actor: CurrentUser) -> None:
    data = store().get(job.object_key)
    header, rows = read_table(data, job.filename)
    profile = job.profile
    cmap: dict[str, str | None] = profile.get("column_map", {})
    smap: dict[str, str] = {k.strip().lower(): v for k, v in (profile.get("status_map") or {}).items() if v}
    catmap: dict[str, str] = {k.strip().lower(): v for k, v in (profile.get("category_map") or {}).items() if v}
    channel = profile.get("channel", "REGISTER")
    fmt = profile.get("date_format")
    idx = {h: i for i, h in enumerate(header)}
    missing = [f for f in REQUIRED if not cmap.get(f) or cmap[f] not in idx]
    if not (cmap.get("institution_name") in idx or cmap.get("permanent_id") in idx):
        missing.append("institution_name or permanent_id")
    job.total_rows = len(rows)
    if missing:
        job.status, job.errors = "FAILED", [{"row": 0, "error": f"Required columns not mapped: {', '.join(missing)}"}]
        job.finished_at = dt.datetime.now(dt.UTC)
        return

    def get(row: list[Any], fld: str) -> Any:
        col = cmap.get(fld)
        if not col or col not in idx or idx[col] >= len(row):
            return None
        v = row[idx[col]]
        return v.strip() if isinstance(v, str) else v

    errors: list[dict] = []
    created = dupes = 0
    for n, row in enumerate(rows, start=2):
        try:
            received = _parse_date(get(row, "received_on"), fmt)
            if received is None:
                raise ValueError(f"Receipt date '{get(row, 'received_on')}' is not a recognised date")
            raw_cat = str(get(row, "category") or "").strip()
            category = catmap.get(raw_cat.lower()) or (raw_cat.upper() if raw_cat else None)
            if not category:
                raise ValueError("Category is empty")
            raw_status = str(get(row, "status") or "").strip()
            status = None
            if raw_status:
                status = smap.get(raw_status.lower()) or (
                    raw_status.upper() if raw_status.upper() in fsm.STATUSES else None
                )
                if status is None:
                    raise ValueError(f"Status '{raw_status}' is not in the status mapping")
                if status == "TRANSFER_PROPOSED":
                    status = "REGISTERED"
            body = str(get(row, "body") or "").strip()
            if len(body) < 10:
                raise ValueError("Complaint details are missing or shorter than 10 characters")
            subject = str(get(row, "subject") or body[:120]).strip()
            ref = get(row, "source_ref")
            pin = get(row, "pin")
            inp = CaseInput(
                channel=channel,
                category_code=category,
                received_on=received,
                subject=subject[:300],
                body=body,
                institution_name=str(get(row, "institution_name") or "") or None,
                institution_city=str(get(row, "city") or "") or None,
                institution_state=get(row, "state"),
                institution_pin=str(pin).split(".")[0] if pin not in (None, "") else None,
                permanent_id=str(get(row, "permanent_id") or "") or None,
                source_ref=str(ref) if ref not in (None, "") else None,
                unit_file_no=str(get(row, "unit_file_no") or "") or None,
                complainant_name=get(row, "complainant_name"),
                complainant_phone=str(get(row, "complainant_phone") or "") or None,
                complainant_email=get(row, "complainant_email"),
                initial_status=status,
                next_action=get(row, "next_action"),
                disposed_on=_parse_date(get(row, "disposed_on"), fmt),
            )
            with db.begin_nested():
                _, was_created = create_case(db, actor, inp, import_job_id=job.id)
            if was_created:
                created += 1
            else:
                dupes += 1
        except AppError as exc:
            errors.append({"row": n, "error": exc.detail or exc.title})
        except (ValueError, TypeError) as exc:
            errors.append({"row": n, "error": str(exc)})
        if n % 200 == 0:
            job.created_rows, job.duplicate_rows, job.error_rows = created, dupes, len(errors)
            db.commit()
    job.created_rows, job.duplicate_rows, job.error_rows = created, dupes, len(errors)
    job.errors = errors[:MAX_ERRORS]
    job.status, job.finished_at = "DONE", dt.datetime.now(dt.UTC)


def new_job(db: Session, actor: CurrentUser, filename: str, sha: str, key: str, profile: dict) -> ImportJob:
    if profile.get("channel") not in CHANNELS:
        raise AppError(422, "Invalid profile", "Choose the channel these records came from.", "invalid")
    job = ImportJob(
        id=uuid.uuid4(),
        unit_id=actor.unit_id,
        uploaded_by=actor.id,
        filename=filename,
        file_sha256=sha,
        object_key=key,
        profile=profile,
    )
    db.add(job)
    db.flush()
    return job
